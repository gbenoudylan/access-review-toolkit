"""
Cohérence transversale : une correction manuelle de statut (mapping appris
dans le dashboard) doit produire la MÊME lecture « actif / inactif » dans
TOUS les modules — analyse, orphelins, partis-mais-actifs, doublons,
transferts, comparaison avec la revue précédente, rapports et contrôle
qualité. Avant correction, chaque module refaisait sa propre lecture du
statut et ignorait les mappings manuels.
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))

import pandas as pd
from analysis.access_review import analyze_access, summarize
from analysis.hr_crossref import flag_transferred_but_still_active
from ingestion.ingest import compute_data_quality_report
from reporting.export import _active_accounts, _compute_comparison_stats

MAPPINGS = {"valid": "active", "suspendu": "inactive"}


def _frame():
    return pd.DataFrame({
        "username": ["admin", "jdupont1", "jdupont2", "mmartin", "old_admin"],
        "full_name": ["", "Jean Dupont", "Jean Dupont", "Marie Martin", ""],
        "system": ["AD"] * 5,
        "account_status": ["Valid", "Valid", "Valid", "Valid", "Suspendu"],
        "employee_status": ["Active", "Active", "Active", "Terminated", "Active"],
        "last_login_date": ["2026-09-01"] * 5,
    })


def test_all_modules_agree_on_active_status_with_manual_mapping():
    df = analyze_access(_frame(), custom_status_mappings=MAPPINGS)

    # 1. Statut résolu : Valid -> actif ; Suspendu -> inactif
    assert df["is_active_for_audit"].tolist() == [True, True, True, True, False]

    # 2. Orphelins : 'admin' (actif) oui ; 'old_admin' (inactif mappé) non
    assert df["is_orphaned_account"].tolist() == [True, False, False, False, False]
    assert summarize(df)["orphaned_accounts"] == 1

    # 3. Parti mais actif : mmartin (Terminated + Valid mappé actif)
    assert df["is_terminated_but_active"].tolist() == [False, False, False, True, False]

    # 4. Doublons actifs : les deux 'Jean Dupont'
    assert df["is_duplicate_account"].tolist() == [False, True, True, False, False]

    # 5. Rapports : la liste « comptes actifs » suit le même statut
    assert len(_active_accounts(df)) == 4

    # 6. Transferts : un transféré au statut « Valid » (mappé actif) est signalé
    transferred = pd.DataFrame({"full_name": ["Marie Martin"],
                                "old_department": ["A"], "new_department": ["B"]})
    flagged = flag_transferred_but_still_active(df.copy(), transferred)
    assert flagged["is_transferred_but_active"].tolist() == [False, False, False, True, False]

    # 7. Contrôle qualité : « Valid » mappé n'est plus une valeur de statut inconnue
    assert compute_data_quality_report(df)["issues"]["unknown_status"] == 0


def test_comparison_reactivated_uses_resolved_status_on_both_cycles():
    prev_raw = pd.DataFrame({
        "username": ["a", "b", "c"], "system": ["AD"] * 3,
        "account_status": ["Suspendu", "Valid", "Suspendu"],
        "last_login_date": ["2026-06-01"] * 3,
    })
    curr_raw = pd.DataFrame({
        "username": ["a", "b", "c"], "system": ["AD"] * 3,
        "account_status": ["Valid", "Valid", "Suspendu"],
        "last_login_date": ["2026-09-01"] * 3,
    })
    prev = analyze_access(prev_raw, custom_status_mappings=MAPPINGS)
    curr = analyze_access(curr_raw, custom_status_mappings=MAPPINGS)
    stats = _compute_comparison_stats(curr, prev)
    # a: inactif -> actif = réactivé ; b: actif -> actif ; c: inactif -> inactif
    assert list(stats["reactivated_accounts"]) == ["a"]


def _password_df():
    rows = [
        ("jdupont", "Active", "2025-01-10"), ("kbrou", "Active", "2025-02-01"),
        ("aeya", "Active", "2026-09-01"),
        ("old1", "Disabled", "2024-01-01"), ("old2", "Disabled", "2024-03-01"),
        ("old3", "Locked", "2023-05-01"),
        ("svc_old", "Active", "2024-06-01"), ("svc_dis", "Disabled", "2023-06-01"),
    ]
    return analyze_access(pd.DataFrame({
        "username": [r[0] for r in rows], "full_name": [r[0].title() for r in rows],
        "system": "AD", "account_status": [r[1] for r in rows],
        "password_last_set": [r[2] for r in rows], "last_login_date": "2026-09-20",
    }))


def _word_summary_and_detail(df, tmp_path):
    import re
    import docx
    from reporting.export import generate_word_report
    d = docx.Document(generate_word_report(
        df, tmp_path / "r.docx", current_extraction_date="2026-09-29", period="T3 2026"))
    summary = {}
    for t in d.tables:
        header = [c.text for c in t.rows[0].cells]
        if header[:2] == ["No.", "Control"] and "Findings" in header:
            for r in t.rows[1:]:
                cells = [x.text for x in r.cells]
                summary[int(cells[0])] = cells[3]
    detail, current = {}, None
    for p in d.paragraphs:
        m = re.match(r"^(\d{1,2})\.\s*\S", p.text)
        if m and len(p.text) < 110 and "account(s)" not in p.text:
            current = int(m.group(1))
        m2 = re.search(r"(\d+) account\(s\) concerned", p.text)
        if m2 and current is not None and current not in detail:
            detail[current] = m2.group(1)
    return summary, detail


def test_summarize_password_stale_counts_active_accounts_only():
    df = _password_df()
    active_stale = int((df["is_password_stale"] & df["is_active_for_audit"]).sum())
    assert summarize(df)["password_stale"] == active_stale
    assert int(df["is_password_stale"].sum()) > active_stale  # des inactifs périmés existent


def test_word_control_summary_matches_section_details(tmp_path=None):
    """Régression : le sommaire du Word recomptait chaque contrôle sur TOUS les
    comptes (inactifs inclus) alors que chaque section détail ne garde que les
    comptes actifs -> ex. Ctrl 14 : 6 au sommaire, 3 dans le détail."""
    import tempfile
    from pathlib import Path
    tmp = Path(tempfile.mkdtemp()) if tmp_path is None else tmp_path
    df = _password_df()
    summary, detail = _word_summary_and_detail(df, tmp)
    expected = int((df["is_password_stale"] & df["is_active_for_audit"]).sum())
    assert summary[14] == str(expected) == detail[14]
    # Ctrl 18 exclu : N/A au sommaire sans source RH, la section détail
    # l'indique elle-même (« No terminated employees list provided »).
    for number in (set(summary) & set(detail)) - {18}:
        assert summary[number] == detail[number], (number, summary[number], detail[number])


def test_excel_summary_shows_comparison_controls_when_previous_review_given():
    import tempfile
    from pathlib import Path
    import openpyxl
    from reporting.export import generate_excel_report
    prev = analyze_access(pd.DataFrame({"username": ["a", "b"], "system": "AD",
                                         "account_status": "Active", "last_login_date": "2026-06-01"}))
    curr = analyze_access(pd.DataFrame({"username": ["a", "b", "c"], "system": "AD",
                                         "account_status": "Active", "last_login_date": "2026-09-01"}))
    out = generate_excel_report(curr, Path(tempfile.mkdtemp()) / "r.xlsx", previous_df=prev,
                                current_extraction_date="2026-09-29", previous_extraction_date="2026-06-30")
    rows = {r[0]: r[3] for r in openpyxl.load_workbook(out)["Summary"].iter_rows(values_only=True)
            if isinstance(r[0], int)}
    assert str(rows[10]) == "1"  # 1 compte créé (c) ; "—" (N/A) avant correctif


def test_report_wording_no_country_and_new_created_expectation():
    from reporting.template_sections import TEMPLATE_CONTROLS, CONTROL_SUBSECTIONS
    blob = repr(TEMPLATE_CONTROLS) + repr(CONTROL_SUBSECTIONS)
    assert "LIBERIA" not in blob.upper()
    created = next(g for n, _, g, _ in CONTROL_SUBSECTIONS if n == 10)
    assert created == (
        "Expectations: All new accounts must be created following an approved request "
        "(e.g. via ServiceNow ticket approved or via user access form approved). "
        "The system owner must share that approved ticket ID or user access form approved "
        "that was used to create the(se) account(s)."
    )


def test_status_interpretation_labels_have_no_emoji():
    """Les libellés de statut affichés dans la zone principale du dashboard
    (écran de configuration) ne contiennent plus d'emoji."""
    import re
    df = analyze_access(pd.DataFrame({
        "username": ["a", "b", "c", "d"], "system": "AD",
        "account_status": ["Active", "Disabled", "Locked", "Valid"],
        "last_login_date": "2026-09-01",
    }))
    labels = df.attrs["all_status_values"]
    assert set(labels) == {"Active", "Disabled", "Locked", "Valid"}
    emoji = re.compile("[\U0001F300-\U0001FAFF\u2600-\u27BF\u2B00-\u2BFF\uFE0F]")
    assert not any(emoji.search(v) for v in labels.values()), labels
    assert labels["Valid"].startswith("Unknown")


if __name__ == "__main__":
    test_status_interpretation_labels_have_no_emoji()
    test_all_modules_agree_on_active_status_with_manual_mapping()
    test_comparison_reactivated_uses_resolved_status_on_both_cycles()
    test_summarize_password_stale_counts_active_accounts_only()
    test_word_control_summary_matches_section_details()
    test_excel_summary_shows_comparison_controls_when_previous_review_given()
    test_report_wording_no_country_and_new_created_expectation()
    print("OK - cohérence transversale")
